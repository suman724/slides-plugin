"""Chart generation helpers for python-pptx."""

from pptx.util import Inches, Pt
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION

from design_system import Theme, Grid


CHART_TYPE_MAP = {
    "column_clustered": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "column_stacked": XL_CHART_TYPE.COLUMN_STACKED,
    "bar_clustered": XL_CHART_TYPE.BAR_CLUSTERED,
    "bar_stacked": XL_CHART_TYPE.BAR_STACKED,
    "line": XL_CHART_TYPE.LINE,
    "line_markers": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
    "area": XL_CHART_TYPE.AREA,
    "area_stacked": XL_CHART_TYPE.AREA_STACKED,
    "doughnut": XL_CHART_TYPE.DOUGHNUT,
}


def add_chart(slide, chart_spec, left, top, width, height, theme):
    """Add a chart to the slide from a chart specification.

    chart_spec: dict with keys:
        - chart_type: str (maps to CHART_TYPE_MAP)
        - categories: list of str
        - series: list of dicts with 'name' and 'values'
    """
    chart_type_str = chart_spec.get("chart_type", "column_clustered")
    xl_chart_type = CHART_TYPE_MAP.get(chart_type_str, XL_CHART_TYPE.COLUMN_CLUSTERED)

    chart_data = CategoryChartData()
    chart_data.categories = chart_spec.get("categories", [])

    for series in chart_spec.get("series", []):
        chart_data.add_series(series["name"], series["values"])

    chart_frame = slide.shapes.add_chart(
        xl_chart_type,
        Inches(left), Inches(top),
        Inches(width), Inches(height),
        chart_data,
    )

    chart = chart_frame.chart
    _style_chart(chart, theme, xl_chart_type)

    return chart_frame


def _style_chart(chart, theme, chart_type):
    """Apply theme-consistent styling to a chart."""
    # Legend
    chart.has_legend = True
    chart.legend.include_in_layout = False
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.font.size = Pt(10)
    chart.legend.font.name = theme.typography.body.name

    # Apply palette colors to series
    plot = chart.plots[0]
    for i, series in enumerate(plot.series):
        color_idx = i % len(theme.chart_palette)
        fill = series.format.fill
        fill.solid()
        fill.fore_color.rgb = theme.chart_palette[color_idx]

        # For line charts, also style the line
        if chart_type in (XL_CHART_TYPE.LINE, XL_CHART_TYPE.LINE_MARKERS):
            series.format.line.color.rgb = theme.chart_palette[color_idx]
            series.format.line.width = Pt(2.5)
            series.smooth = False

    # Category axis styling
    if hasattr(chart, 'category_axis'):
        cat_axis = chart.category_axis
        cat_axis.tick_labels.font.size = Pt(9)
        cat_axis.tick_labels.font.name = theme.typography.body.name
        cat_axis.has_major_gridlines = False

    # Value axis styling
    if hasattr(chart, 'value_axis'):
        val_axis = chart.value_axis
        val_axis.tick_labels.font.size = Pt(9)
        val_axis.tick_labels.font.name = theme.typography.body.name
        val_axis.has_major_gridlines = True
        val_axis.major_gridlines.format.line.color.rgb = theme.background_alt

    # Pie charts: no axes, add data labels
    if chart_type in (XL_CHART_TYPE.PIE, XL_CHART_TYPE.DOUGHNUT):
        plot.has_data_labels = True
        data_labels = plot.data_labels
        data_labels.font.size = Pt(10)
        data_labels.font.name = theme.typography.body.name
        data_labels.show_percentage = True
        data_labels.show_category_name = True
        data_labels.show_value = False
