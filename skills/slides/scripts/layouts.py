"""Slide layout rendering functions and layout registry.

Each layout function has the signature:
    render_<type>_slide(slide, spec, theme, grid)

All layouts use blank slides (Layout 6) and build everything from shapes.
"""

from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

from design_system import Theme, Grid, FontSpec
from components import (
    add_solid_bg, add_gradient_bg, add_shape_bg,
    add_text_box, add_slide_title, add_accent_bar,
    add_bullet_list, add_metric_card, add_numbered_item,
    add_timeline_marker, add_pro_con_list, add_rounded_card,
    add_separator_line,
)
from charts import add_chart
from tables import add_table


# ---------------------------------------------------------------------------
# Title Slide
# ---------------------------------------------------------------------------

def render_title_slide(slide, spec, theme, grid):
    """Cover slide with gradient background, title, subtitle, and footer info."""
    add_gradient_bg(slide, theme.primary, theme.secondary)

    # Accent bar (left side vertical stripe)
    add_shape_bg(slide, 0.60, 1.8, 0.08, 3.5, theme.accent)

    # Title
    add_text_box(
        slide, spec.get("title", ""),
        left=1.0, top=2.0,
        width=grid.content_width - 0.5, height=1.2,
        font_spec=theme.typography.title_slide_title,
        color=theme.text_on_primary,
    )

    # Subtitle
    subtitle = spec.get("subtitle", "")
    if subtitle:
        add_text_box(
            slide, subtitle,
            left=1.0, top=3.3,
            width=grid.content_width - 0.5, height=0.60,
            font_spec=theme.typography.title_slide_subtitle,
            color=theme.text_on_primary,
        )

    # Footer band
    add_shape_bg(slide, 0, grid.slide_height - 0.80, grid.slide_width, 0.80, theme.primary)

    # Footer text: date left, company right
    date_text = spec.get("date", "")
    if not date_text:
        # fallback from metadata is handled at the generate_pptx level
        pass
    if date_text:
        add_text_box(
            slide, date_text,
            left=grid.margin_left, top=grid.slide_height - 0.65,
            width=4.0, height=0.40,
            font_spec=theme.typography.footer,
            color=theme.text_on_primary,
        )

    footer_right = spec.get("footer_right", spec.get("confidentiality", ""))
    if footer_right:
        add_text_box(
            slide, footer_right,
            left=grid.slide_width - grid.margin_right - 4.0,
            top=grid.slide_height - 0.65,
            width=4.0, height=0.40,
            font_spec=theme.typography.footer,
            color=theme.text_on_primary,
            alignment=PP_ALIGN.RIGHT,
        )


# ---------------------------------------------------------------------------
# Closing Slide
# ---------------------------------------------------------------------------

def render_closing_slide(slide, spec, theme, grid):
    """Thank you / closing slide."""
    add_gradient_bg(slide, theme.primary, theme.secondary)

    # Title centered
    add_text_box(
        slide, spec.get("title", "Thank You"),
        left=1.0, top=2.2,
        width=grid.content_width - 0.5, height=1.0,
        font_spec=FontSpec("Calibri", 44, bold=True),
        color=theme.text_on_primary,
        alignment=PP_ALIGN.CENTER,
    )

    # Subtitle
    subtitle = spec.get("subtitle", "")
    if subtitle:
        add_text_box(
            slide, subtitle,
            left=1.0, top=3.4,
            width=grid.content_width - 0.5, height=0.60,
            font_spec=theme.typography.title_slide_subtitle,
            color=theme.text_on_primary,
            alignment=PP_ALIGN.CENTER,
        )

    # Contact
    contact = spec.get("contact", "")
    if contact:
        add_text_box(
            slide, contact,
            left=1.0, top=4.3,
            width=grid.content_width - 0.5, height=0.40,
            font_spec=theme.typography.body,
            color=theme.text_on_primary,
            alignment=PP_ALIGN.CENTER,
        )


# ---------------------------------------------------------------------------
# Agenda Slide
# ---------------------------------------------------------------------------

def render_agenda_slide(slide, spec, theme, grid):
    """Table of contents with numbered items."""
    add_slide_title(slide, spec.get("title", "Agenda"), theme, grid)

    items = spec.get("items", [])
    item_height = 0.75
    start_top = grid.content_top + 0.15

    for i, item in enumerate(items):
        number = item.get("number", i + 1)
        text = item.get("text", "")
        subtext = item.get("subtext", "")

        add_numbered_item(
            slide, number, text, subtext,
            left=grid.margin_left + 0.30,
            top=start_top + i * item_height,
            width=grid.content_width - 0.60,
            height=item_height,
            theme=theme,
        )


# ---------------------------------------------------------------------------
# Section Divider
# ---------------------------------------------------------------------------

def render_section_divider(slide, spec, theme, grid):
    """Section break slide with large number and title."""
    # Left band with primary color
    add_shape_bg(slide, 0, 0, 4.5, grid.slide_height, theme.primary)

    # Section number in left band
    section_num = spec.get("section_number", "")
    if section_num:
        add_text_box(
            slide, str(section_num).zfill(2),
            left=0.5, top=2.0,
            width=3.5, height=1.5,
            font_spec=FontSpec("Calibri", 72, bold=True),
            color=theme.text_on_primary,
            alignment=PP_ALIGN.CENTER,
        )

    # Title on right
    add_text_box(
        slide, spec.get("title", ""),
        left=5.2, top=2.2,
        width=grid.slide_width - 5.2 - grid.margin_right, height=1.0,
        font_spec=theme.typography.section_title,
        color=theme.text_primary,
    )

    # Accent bar under title
    add_accent_bar(
        slide, left=5.2, top=3.3,
        width=2.0, theme=theme,
    )

    # Subtitle
    subtitle = spec.get("subtitle", "")
    if subtitle:
        add_text_box(
            slide, subtitle,
            left=5.2, top=3.55,
            width=grid.slide_width - 5.2 - grid.margin_right, height=0.60,
            font_spec=theme.typography.body,
            color=theme.text_secondary,
        )


# ---------------------------------------------------------------------------
# Content Slide (standard bullets)
# ---------------------------------------------------------------------------

def render_content_slide(slide, spec, theme, grid):
    """Standard slide with title and bullet list."""
    add_slide_title(slide, spec.get("title", ""), theme, grid)

    body = spec.get("body", [])
    if body:
        add_bullet_list(
            slide, body,
            left=grid.margin_left + 0.20,
            top=grid.content_top + 0.15,
            width=grid.content_width - 0.40,
            height=grid.content_height - 0.30,
            theme=theme, grid=grid,
        )


# ---------------------------------------------------------------------------
# Two-Column Slide
# ---------------------------------------------------------------------------

def render_two_column_slide(slide, spec, theme, grid):
    """Side-by-side two-column layout."""
    add_slide_title(slide, spec.get("title", ""), theme, grid)

    col_width = grid.col_width(2)
    col_top = grid.content_top + 0.15

    for col_idx, side in enumerate(["left", "right"]):
        col_data = spec.get(side, {})
        col_left = grid.col_left(col_idx, 2)

        # Column heading
        heading = col_data.get("heading", "")
        if heading:
            add_text_box(
                slide, heading,
                left=col_left, top=col_top,
                width=col_width, height=0.40,
                font_spec=theme.typography.heading,
                color=theme.primary,
            )
            # Thin separator under heading
            add_separator_line(
                slide, left=col_left, top=col_top + 0.42,
                width=col_width, theme=theme,
            )

        # Column body
        body = col_data.get("body", [])
        if body:
            add_bullet_list(
                slide, body,
                left=col_left + 0.10,
                top=col_top + 0.55,
                width=col_width - 0.20,
                height=grid.content_height - 0.80,
                theme=theme, grid=grid,
            )


# ---------------------------------------------------------------------------
# Three-Column Slide
# ---------------------------------------------------------------------------

def render_three_column_slide(slide, spec, theme, grid):
    """Triple-column layout."""
    add_slide_title(slide, spec.get("title", ""), theme, grid)

    col_width = grid.col_width(3)
    col_top = grid.content_top + 0.15

    columns = spec.get("columns", [])
    for col_idx, col_data in enumerate(columns[:3]):
        col_left = grid.col_left(col_idx, 3)

        heading = col_data.get("heading", "")
        if heading:
            add_text_box(
                slide, heading,
                left=col_left, top=col_top,
                width=col_width, height=0.40,
                font_spec=theme.typography.heading,
                color=theme.primary,
            )
            add_separator_line(
                slide, left=col_left, top=col_top + 0.42,
                width=col_width, theme=theme,
            )

        body = col_data.get("body", [])
        if body:
            add_bullet_list(
                slide, body,
                left=col_left + 0.10,
                top=col_top + 0.55,
                width=col_width - 0.20,
                height=grid.content_height - 0.80,
                theme=theme, grid=grid,
            )


# ---------------------------------------------------------------------------
# Chart Slide
# ---------------------------------------------------------------------------

def render_chart_slide(slide, spec, theme, grid):
    """Data visualization slide with chart and optional annotation."""
    add_slide_title(slide, spec.get("title", ""), theme, grid)

    annotation = spec.get("annotation", "")
    chart_height = grid.content_height - 0.60 if annotation else grid.content_height - 0.20

    add_chart(
        slide, spec,
        left=grid.margin_left + 0.30,
        top=grid.content_top + 0.10,
        width=grid.content_width - 0.60,
        height=chart_height,
        theme=theme,
    )

    if annotation:
        add_text_box(
            slide, annotation,
            left=grid.margin_left + 0.30,
            top=grid.content_top + chart_height + 0.15,
            width=grid.content_width - 0.60,
            height=0.35,
            font_spec=theme.typography.caption,
            color=theme.text_secondary,
            alignment=PP_ALIGN.LEFT,
        )


# ---------------------------------------------------------------------------
# Table Slide
# ---------------------------------------------------------------------------

def render_table_slide(slide, spec, theme, grid):
    """Data table slide."""
    add_slide_title(slide, spec.get("title", ""), theme, grid)

    add_table(
        slide, spec,
        left=grid.margin_left + 0.20,
        top=grid.content_top + 0.15,
        width=grid.content_width - 0.40,
        height=grid.content_height - 0.30,
        theme=theme,
    )


# ---------------------------------------------------------------------------
# Key Metrics (KPI Dashboard)
# ---------------------------------------------------------------------------

def render_key_metrics_slide(slide, spec, theme, grid):
    """KPI dashboard with 3-4 metric cards in a row."""
    add_slide_title(slide, spec.get("title", ""), theme, grid)

    metrics = spec.get("metrics", [])
    n = len(metrics)
    if n == 0:
        return

    n = min(n, 4)  # max 4 cards
    card_gutter = 0.35
    total_gutter = (n - 1) * card_gutter
    card_width = (grid.content_width - total_gutter) / n
    card_height = 1.90
    card_top = grid.content_top + (grid.content_height - card_height) / 2

    for i, metric in enumerate(metrics[:n]):
        card_left = grid.margin_left + i * (card_width + card_gutter)
        add_metric_card(
            slide, metric,
            left=card_left, top=card_top,
            width=card_width, height=card_height,
            theme=theme,
        )


# ---------------------------------------------------------------------------
# Comparison Slide (Pros/Cons)
# ---------------------------------------------------------------------------

def render_comparison_slide(slide, spec, theme, grid):
    """Side-by-side comparison with pros/cons lists."""
    add_slide_title(slide, spec.get("title", ""), theme, grid)

    col_width = grid.col_width(2)
    col_top = grid.content_top + 0.10

    for col_idx, side in enumerate(["left", "right"]):
        col_data = spec.get(side, {})
        col_left = grid.col_left(col_idx, 2)

        # Card background
        add_rounded_card(
            slide, col_left, col_top,
            col_width, grid.content_height - 0.20,
            theme.background_alt,
        )

        # Heading
        heading = col_data.get("heading", "")
        if heading:
            add_text_box(
                slide, heading,
                left=col_left + 0.20, top=col_top + 0.15,
                width=col_width - 0.40, height=0.40,
                font_spec=theme.typography.heading,
                color=theme.primary,
            )
            add_accent_bar(
                slide, left=col_left + 0.20, top=col_top + 0.58,
                width=col_width - 0.40, theme=theme,
            )

        # Pros
        pros = col_data.get("pros", [])
        pros_top = col_top + 0.75
        if pros:
            add_text_box(
                slide, "Advantages",
                left=col_left + 0.20, top=pros_top,
                width=col_width - 0.40, height=0.30,
                font_spec=FontSpec("Calibri", 11, bold=True),
                color=theme.positive,
            )
            add_pro_con_list(
                slide, pros, is_pro=True,
                left=col_left + 0.20, top=pros_top + 0.30,
                width=col_width - 0.40, theme=theme,
            )

        # Cons
        cons = col_data.get("cons", [])
        cons_top = pros_top + 0.30 + len(pros) * 0.32 + 0.20
        if cons:
            add_text_box(
                slide, "Challenges",
                left=col_left + 0.20, top=cons_top,
                width=col_width - 0.40, height=0.30,
                font_spec=FontSpec("Calibri", 11, bold=True),
                color=theme.negative,
            )
            add_pro_con_list(
                slide, cons, is_pro=False,
                left=col_left + 0.20, top=cons_top + 0.30,
                width=col_width - 0.40, theme=theme,
            )


# ---------------------------------------------------------------------------
# Timeline Slide
# ---------------------------------------------------------------------------

def render_timeline_slide(slide, spec, theme, grid):
    """Horizontal timeline with milestone markers."""
    add_slide_title(slide, spec.get("title", ""), theme, grid)

    milestones = spec.get("milestones", [])
    n = len(milestones)
    if n == 0:
        return

    # Horizontal connector line
    line_top = grid.content_top + grid.content_height * 0.40
    line_left = grid.margin_left + 0.50
    line_width = grid.content_width - 1.00

    add_shape_bg(
        slide, line_left, line_top + 0.13,
        line_width, 0.04, theme.text_secondary,
    )

    # Place milestones evenly along the line
    spacing = line_width / max(n - 1, 1) if n > 1 else 0
    for i, milestone in enumerate(milestones):
        center_x = line_left + i * spacing if n > 1 else line_left + line_width / 2
        add_timeline_marker(
            slide, milestone, center_x, line_top,
            theme=theme,
            is_first=(i == 0),
            is_last=(i == n - 1),
        )


# ---------------------------------------------------------------------------
# Takeaway / Summary Slide
# ---------------------------------------------------------------------------

def render_takeaway_slide(slide, spec, theme, grid):
    """Key findings with numbered points and optional call-to-action."""
    add_slide_title(slide, spec.get("title", "Key Takeaways"), theme, grid)

    points = spec.get("points", [])
    item_height = 0.65
    start_top = grid.content_top + 0.15

    for i, point in enumerate(points):
        add_numbered_item(
            slide, i + 1, point, "",
            left=grid.margin_left + 0.30,
            top=start_top + i * item_height,
            width=grid.content_width - 0.60,
            height=item_height,
            theme=theme,
        )

    # Call to action box
    cta = spec.get("call_to_action", "")
    if cta:
        cta_top = start_top + len(points) * item_height + 0.30
        add_rounded_card(
            slide,
            grid.margin_left + 0.30, cta_top,
            grid.content_width - 0.60, 0.60,
            theme.accent,
        )
        add_text_box(
            slide, cta,
            left=grid.margin_left + 0.50,
            top=cta_top + 0.10,
            width=grid.content_width - 1.00,
            height=0.40,
            font_spec=FontSpec("Calibri", 14, bold=True),
            color=theme.text_on_primary,
            alignment=PP_ALIGN.CENTER,
        )


# ---------------------------------------------------------------------------
# Quote Slide
# ---------------------------------------------------------------------------

def render_quote_slide(slide, spec, theme, grid):
    """Large quote with attribution."""
    add_slide_title(slide, spec.get("title", ""), theme, grid)

    quote_text = spec.get("quote", "")
    attribution = spec.get("attribution", "")

    # Large opening quote mark
    add_text_box(
        slide, "\u201C",
        left=grid.margin_left + 0.30,
        top=grid.content_top,
        width=1.0, height=1.0,
        font_spec=FontSpec("Calibri", 72, bold=True),
        color=theme.accent,
    )

    # Quote text
    add_text_box(
        slide, quote_text,
        left=grid.margin_left + 1.2,
        top=grid.content_top + 0.40,
        width=grid.content_width - 2.0,
        height=grid.content_height - 1.50,
        font_spec=FontSpec("Calibri", 20, italic=True),
        color=theme.text_primary,
    )

    # Attribution
    if attribution:
        add_text_box(
            slide, f"\u2014 {attribution}",
            left=grid.margin_left + 1.2,
            top=grid.content_bottom - 1.0,
            width=grid.content_width - 2.0,
            height=0.40,
            font_spec=theme.typography.body,
            color=theme.text_secondary,
            alignment=PP_ALIGN.RIGHT,
        )


# ---------------------------------------------------------------------------
# Image with Text Slide
# ---------------------------------------------------------------------------

def render_image_with_text_slide(slide, spec, theme, grid):
    """Image placeholder on left, text content on right."""
    add_slide_title(slide, spec.get("title", ""), theme, grid)

    col_width = grid.col_width(2)

    # Left: image placeholder (gray box with text)
    left_col = grid.col_left(0, 2)
    img_top = grid.content_top + 0.15
    img_height = grid.content_height - 0.30

    image_path = spec.get("image_path", "")
    if image_path:
        try:
            slide.shapes.add_picture(
                image_path,
                Inches(left_col), Inches(img_top),
                Inches(col_width), Inches(img_height),
            )
        except Exception:
            _add_image_placeholder(slide, left_col, img_top, col_width, img_height, theme)
    else:
        _add_image_placeholder(slide, left_col, img_top, col_width, img_height, theme)

    # Right: text content
    right_col = grid.col_left(1, 2)
    body = spec.get("body", [])
    if body:
        add_bullet_list(
            slide, body,
            left=right_col + 0.10,
            top=grid.content_top + 0.15,
            width=col_width - 0.20,
            height=grid.content_height - 0.30,
            theme=theme, grid=grid,
        )
    else:
        text = spec.get("text", "")
        if text:
            add_text_box(
                slide, text,
                left=right_col + 0.10,
                top=grid.content_top + 0.15,
                width=col_width - 0.20,
                height=grid.content_height - 0.30,
                font_spec=theme.typography.body,
                color=theme.text_primary,
            )


def _add_image_placeholder(slide, left, top, width, height, theme):
    """Add a gray placeholder rectangle with 'Image' text."""
    add_rounded_card(slide, left, top, width, height, theme.background_alt)
    add_text_box(
        slide, "[Image]",
        left=left, top=top + height / 2 - 0.25,
        width=width, height=0.50,
        font_spec=FontSpec("Calibri", 18),
        color=theme.text_secondary,
        alignment=PP_ALIGN.CENTER,
    )


# ---------------------------------------------------------------------------
# Big Number Slide
# ---------------------------------------------------------------------------

def render_big_number_slide(slide, spec, theme, grid):
    """Single standout statistic with context text."""
    add_slide_title(slide, spec.get("title", ""), theme, grid)

    value = spec.get("value", "")
    label = spec.get("label", "")
    context = spec.get("context", "")

    # Large number centered
    add_text_box(
        slide, value,
        left=grid.margin_left,
        top=grid.content_top + 0.30,
        width=grid.content_width,
        height=1.5,
        font_spec=FontSpec("Calibri", 80, bold=True),
        color=theme.accent,
        alignment=PP_ALIGN.CENTER,
    )

    # Label
    if label:
        add_text_box(
            slide, label,
            left=grid.margin_left,
            top=grid.content_top + 1.90,
            width=grid.content_width,
            height=0.60,
            font_spec=theme.typography.heading,
            color=theme.text_primary,
            alignment=PP_ALIGN.CENTER,
        )

    # Context text
    if context:
        add_text_box(
            slide, context,
            left=grid.margin_left + 1.5,
            top=grid.content_top + 2.70,
            width=grid.content_width - 3.0,
            height=0.80,
            font_spec=theme.typography.body,
            color=theme.text_secondary,
            alignment=PP_ALIGN.CENTER,
        )


# ---------------------------------------------------------------------------
# Layout Registry
# ---------------------------------------------------------------------------

LAYOUT_REGISTRY = {
    "title": render_title_slide,
    "closing": render_closing_slide,
    "agenda": render_agenda_slide,
    "section_divider": render_section_divider,
    "content": render_content_slide,
    "two_column": render_two_column_slide,
    "three_column": render_three_column_slide,
    "chart": render_chart_slide,
    "table": render_table_slide,
    "key_metrics": render_key_metrics_slide,
    "comparison": render_comparison_slide,
    "timeline": render_timeline_slide,
    "takeaway": render_takeaway_slide,
    "quote": render_quote_slide,
    "image_with_text": render_image_with_text_slide,
    "big_number": render_big_number_slide,
}
