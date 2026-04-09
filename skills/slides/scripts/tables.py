"""Table generation helpers for python-pptx."""

from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

from design_system import Theme, Grid
from utils import lighten_color


def add_table(slide, table_spec, left, top, width, height, theme):
    """Add a formatted data table to the slide.

    table_spec: dict with keys:
        - columns: list of str (header labels)
        - rows: list of lists (cell values)
        - highlight_rules: optional list of dicts with column, condition, color
    """
    columns = table_spec.get("columns", [])
    rows = table_spec.get("rows", [])
    highlight_rules = table_spec.get("highlight_rules", [])

    n_rows = len(rows) + 1  # +1 for header
    n_cols = len(columns)

    if n_cols == 0 or n_rows <= 1:
        return None

    table_shape = slide.shapes.add_table(
        n_rows, n_cols,
        Inches(left), Inches(top),
        Inches(width), Inches(height),
    )
    table = table_shape.table

    # Column widths: distribute evenly
    col_width = Inches(width / n_cols)
    for i in range(n_cols):
        table.columns[i].width = col_width

    # Style header row
    for col_idx, col_name in enumerate(columns):
        cell = table.cell(0, col_idx)
        cell.text = col_name
        _style_cell(cell, theme, is_header=True)

    # Style data rows
    for row_idx, row_data in enumerate(rows):
        is_alt = (row_idx % 2 == 1)
        for col_idx, cell_value in enumerate(row_data):
            cell = table.cell(row_idx + 1, col_idx)
            cell.text = str(cell_value)

            # Check highlight rules
            highlight_color = _get_highlight_color(
                cell_value, col_idx, highlight_rules, theme
            )

            _style_cell(
                cell, theme,
                is_header=False,
                is_alt=is_alt,
                highlight_color=highlight_color,
            )

    return table_shape


def _style_cell(cell, theme, is_header=False, is_alt=False, highlight_color=None):
    """Apply consistent styling to a table cell."""
    tf = cell.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE

    for paragraph in tf.paragraphs:
        paragraph.alignment = PP_ALIGN.CENTER
        for run in paragraph.runs:
            if is_header:
                run.font.name = theme.typography.body.name
                run.font.size = Pt(11)
                run.font.bold = True
                run.font.color.rgb = theme.text_on_primary
            else:
                run.font.name = theme.typography.body.name
                run.font.size = Pt(10)
                run.font.bold = False
                if highlight_color:
                    run.font.color.rgb = highlight_color
                    run.font.bold = True
                else:
                    run.font.color.rgb = theme.text_primary

    # Cell fill
    cell_fill = cell.fill
    cell_fill.solid()
    if is_header:
        cell_fill.fore_color.rgb = theme.primary
    elif is_alt:
        cell_fill.fore_color.rgb = theme.background_alt
    else:
        cell_fill.fore_color.rgb = theme.background

    # Cell margins
    cell.margin_left = Inches(0.08)
    cell.margin_right = Inches(0.08)
    cell.margin_top = Inches(0.04)
    cell.margin_bottom = Inches(0.04)


def _get_highlight_color(cell_value, col_idx, highlight_rules, theme):
    """Determine if a cell should be highlighted based on rules."""
    for rule in highlight_rules:
        if rule.get("column") != col_idx:
            continue

        condition = rule.get("condition", "")
        rule_color = rule.get("color", "")

        value_str = str(cell_value).strip()

        if condition == "negative" and _is_negative(value_str):
            return theme.negative if rule_color == "red" else _resolve_color(rule_color, theme)
        elif condition == "positive" and _is_positive(value_str):
            return theme.positive if rule_color == "green" else _resolve_color(rule_color, theme)

    return None


def _is_negative(value_str):
    """Check if a string represents a negative value."""
    return value_str.startswith("-") or value_str.startswith("(")


def _is_positive(value_str):
    """Check if a string represents a positive value."""
    if not value_str:
        return False
    return value_str.startswith("+") or (value_str[0].isdigit() and not value_str.startswith("0"))


def _resolve_color(color_name, theme):
    """Resolve a named color to an RGBColor from the theme."""
    color_map = {
        "red": theme.negative,
        "green": theme.positive,
        "blue": theme.accent,
        "gray": theme.text_secondary,
    }
    return color_map.get(color_name, theme.text_primary)
