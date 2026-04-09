"""Low-level shape rendering for PPTX slides.

All functions take absolute positions (inches) and a ResolvedStyle.
This module is the only one that touches the python-pptx API directly.
"""

from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _apply_font(run, font_role, color=None):
    """Apply a FontRole spec to a text run."""
    run.font.name = font_role.name
    run.font.size = font_role.pt
    run.font.bold = font_role.bold
    run.font.italic = font_role.italic
    if color:
        run.font.color.rgb = color


# ---------------------------------------------------------------------------
# Background
# ---------------------------------------------------------------------------

def render_solid_bg(slide, color):
    """Fill slide background with a solid color."""
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def render_gradient_bg(slide, color1, color2, slide_w, slide_h):
    """Render a gradient background using a full-slide rectangle."""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0), Inches(0),
        Inches(slide_w), Inches(slide_h),
    )
    shape.line.fill.background()
    fill = shape.fill
    fill.gradient()
    fill.gradient_stops[0].color.rgb = color1
    fill.gradient_stops[0].position = 0.0
    fill.gradient_stops[1].color.rgb = color2
    fill.gradient_stops[1].position = 1.0

    # Send to back
    sp = shape._element
    sp.getparent().remove(sp)
    slide.shapes._spTree.insert(2, sp)


# ---------------------------------------------------------------------------
# Text rendering
# ---------------------------------------------------------------------------

def render_text(slide, text, left, top, width, height, font_role, color=None,
                alignment=PP_ALIGN.LEFT, word_wrap=True):
    """Render a styled text box at absolute position."""
    txbox = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height),
    )
    tf = txbox.text_frame
    tf.word_wrap = word_wrap
    p = tf.paragraphs[0]
    p.alignment = alignment
    run = p.add_run()
    run.text = text
    _apply_font(run, font_role, color)
    return txbox


def render_title(slide, text, bounds, style):
    """Render a slide title with accent bar."""
    left, top, width, height = bounds
    txbox = render_text(
        slide, text, left, top, width, height,
        style.typography.title, style.text_primary,
    )
    txbox.text_frame.vertical_anchor = MSO_ANCHOR.BOTTOM

    # Accent bar below title
    bar_top = top + height + 0.05
    render_accent_bar(slide, left, bar_top, min(1.5, width * 0.15), style)
    return txbox


def render_subtitle(slide, text, bounds, style):
    """Render a subtitle text."""
    left, top, width, height = bounds
    return render_text(
        slide, text, left, top, width, height,
        style.typography.body, style.text_secondary,
    )


# ---------------------------------------------------------------------------
# Decorative elements
# ---------------------------------------------------------------------------

def render_accent_bar(slide, left, top, width, style, height=0.04):
    """Thin accent-colored bar."""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(left), Inches(top), Inches(width), Inches(height),
    )
    shape.line.fill.background()
    shape.fill.solid()
    shape.fill.fore_color.rgb = style.accent
    return shape


def render_separator_line(slide, left, top, width, style):
    """Thin horizontal separator."""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(left), Inches(top), Inches(width), Pt(0.5),
    )
    shape.line.fill.background()
    shape.fill.solid()
    shape.fill.fore_color.rgb = style.text_secondary
    return shape


def render_panel(slide, left, top, width, height, color):
    """Solid-fill rectangle panel."""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(left), Inches(top), Inches(width), Inches(height),
    )
    shape.line.fill.background()
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    return shape


def render_rounded_card(slide, left, top, width, height, fill_color, border_color=None):
    """Rounded rectangle card."""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left), Inches(top), Inches(width), Inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape


# ---------------------------------------------------------------------------
# Bullet list
# ---------------------------------------------------------------------------

def render_bullets(slide, points, left, top, width, height, style):
    """Render a bullet list from semantic points.

    points: list of dicts with 'text' and optional 'subpoints' list.
    """
    txbox = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height),
    )
    tf = txbox.text_frame
    tf.word_wrap = True

    first = True
    for point in points:
        text = point.get("text", point) if isinstance(point, dict) else str(point)
        subpoints = point.get("subpoints", []) if isinstance(point, dict) else []

        # Main bullet
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(6)
        p.space_before = Pt(2)
        run = p.add_run()
        run.text = f"\u2022  {text}"
        _apply_font(run, style.typography.body, style.text_primary)

        # Sub-bullets
        for sub in subpoints:
            sp = tf.add_paragraph()
            sp.space_after = Pt(4)
            sp.level = 1
            srun = sp.add_run()
            srun.text = f"  \u2013  {sub}"
            _apply_font(srun, style.typography.body_small, style.text_secondary)

    return txbox


# ---------------------------------------------------------------------------
# Columns (compare, categorize)
# ---------------------------------------------------------------------------

def render_columns(slide, sides, left, top, width, height, style):
    """Render side-by-side columns with headings and bullets."""
    n = len(sides)
    if n == 0:
        return

    gutter = style.gutter_pct * style.slide_width_inches
    col_width = (width - (n - 1) * gutter) / n

    for i, side in enumerate(sides):
        col_left = left + i * (col_width + gutter)
        heading = side.get("heading", "")
        points = side.get("points", side.get("body", []))

        y = top
        if heading:
            render_text(
                slide, heading, col_left, y, col_width, 0.40,
                style.typography.heading, style.primary,
            )
            render_separator_line(slide, col_left, y + 0.42, col_width, style)
            y += 0.55

        if points:
            # Convert simple strings to point dicts
            point_list = [{"text": p} if isinstance(p, str) else p for p in points]
            render_bullets(slide, point_list, col_left + 0.05, y, col_width - 0.10, height - (y - top), style)


# ---------------------------------------------------------------------------
# Metrics (KPI cards)
# ---------------------------------------------------------------------------

def render_metrics(slide, metrics, left, top, width, height, style):
    """Render KPI metric cards in a row."""
    n = min(len(metrics), 4)
    if n == 0:
        return

    gutter = 0.35
    card_width = (width - (n - 1) * gutter) / n
    card_height = min(height, 1.90)
    card_top = top + (height - card_height) / 2

    for i, metric in enumerate(metrics[:n]):
        card_left = left + i * (card_width + gutter)

        # Card background
        render_rounded_card(slide, card_left, card_top, card_width, card_height, style.background_alt)

        # Value
        render_text(
            slide, metric.get("value", ""),
            card_left + 0.15, card_top + 0.25, card_width - 0.30, 0.70,
            style.typography.metric, style.text_primary,
            alignment=PP_ALIGN.CENTER,
        )

        # Label
        render_text(
            slide, metric.get("label", ""),
            card_left + 0.15, card_top + 0.95, card_width - 0.30, 0.35,
            style.typography.metric_label, style.text_secondary,
            alignment=PP_ALIGN.CENTER,
        )

        # Delta with direction
        change = metric.get("change", metric.get("delta", ""))
        trend = metric.get("trend", metric.get("direction", "up"))
        if change:
            arrow = "\u25B2" if trend == "up" else "\u25BC"
            color = style.positive if trend == "up" else style.negative
            render_text(
                slide, f"{arrow} {change}",
                card_left + 0.15, card_top + 1.30, card_width - 0.30, 0.30,
                style.typography.metric_delta, color,
                alignment=PP_ALIGN.CENTER,
            )


# ---------------------------------------------------------------------------
# Chart
# ---------------------------------------------------------------------------

CHART_TYPE_MAP = {
    "column_clustered": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "column_stacked": XL_CHART_TYPE.COLUMN_STACKED,
    "bar_clustered": XL_CHART_TYPE.BAR_CLUSTERED,
    "bar_stacked": XL_CHART_TYPE.BAR_STACKED,
    "line": XL_CHART_TYPE.LINE,
    "line_markers": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
    "area": XL_CHART_TYPE.AREA,
    "area_stacked": XL_CHART_TYPE.AREA_STACKED,
    "doughnut": XL_CHART_TYPE.DOUGHNUT,
}


def render_chart(slide, chart_spec, left, top, width, height, style):
    """Render a chart from a chart specification."""
    series_list = chart_spec.get("series", [])
    categories = chart_spec.get("categories", [])
    if not series_list or not categories:
        # Nothing to chart -- render a placeholder text instead
        render_text(slide, "[No chart data]", left, top + height / 2 - 0.25, width, 0.50,
                    style.typography.body, style.text_secondary, alignment=PP_ALIGN.CENTER)
        return None

    chart_type_str = chart_spec.get("type", "column_clustered")
    xl_type = CHART_TYPE_MAP.get(chart_type_str, XL_CHART_TYPE.COLUMN_CLUSTERED)

    chart_data = CategoryChartData()
    chart_data.categories = categories
    for series in series_list:
        chart_data.add_series(series["name"], series["values"])

    chart_frame = slide.shapes.add_chart(
        xl_type, Inches(left), Inches(top), Inches(width), Inches(height), chart_data,
    )
    chart = chart_frame.chart

    # Style
    chart.has_legend = True
    chart.legend.include_in_layout = False
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.font.size = Pt(10)
    chart.legend.font.name = style.typography.body.name

    plot = chart.plots[0]
    for i, series in enumerate(plot.series):
        idx = i % len(style.chart_colors)
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = style.chart_colors[idx]
        if xl_type in (XL_CHART_TYPE.LINE, XL_CHART_TYPE.LINE_MARKERS):
            series.format.line.color.rgb = style.chart_colors[idx]
            series.format.line.width = Pt(2.5)

    if hasattr(chart, 'category_axis'):
        chart.category_axis.tick_labels.font.size = Pt(9)
        chart.category_axis.tick_labels.font.name = style.typography.body.name
        chart.category_axis.has_major_gridlines = False

    if hasattr(chart, 'value_axis'):
        chart.value_axis.tick_labels.font.size = Pt(9)
        chart.value_axis.tick_labels.font.name = style.typography.body.name
        chart.value_axis.has_major_gridlines = True

    return chart_frame


# ---------------------------------------------------------------------------
# Table
# ---------------------------------------------------------------------------

def render_table(slide, table_spec, left, top, width, height, style):
    """Render a formatted data table."""
    columns = table_spec.get("columns", [])
    rows = table_spec.get("rows", [])
    if not columns:
        return None

    n_rows = len(rows) + 1
    n_cols = len(columns)

    table_shape = slide.shapes.add_table(
        n_rows, n_cols, Inches(left), Inches(top), Inches(width), Inches(height),
    )
    table = table_shape.table

    col_width = Inches(width / n_cols)
    for i in range(n_cols):
        table.columns[i].width = col_width

    # Header row
    for ci, col_name in enumerate(columns):
        cell = table.cell(0, ci)
        cell.text = col_name
        _style_table_cell(cell, style, is_header=True)

    # Data rows
    highlight_rules = table_spec.get("highlight_rules", [])
    for ri, row_data in enumerate(rows):
        for ci, val in enumerate(row_data):
            cell = table.cell(ri + 1, ci)
            cell.text = str(val)
            highlight = _get_highlight(val, ci, highlight_rules, style)
            _style_table_cell(cell, style, is_alt=(ri % 2 == 1), highlight_color=highlight)

    return table_shape


def _style_table_cell(cell, style, is_header=False, is_alt=False, highlight_color=None):
    tf = cell.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for p in tf.paragraphs:
        p.alignment = PP_ALIGN.CENTER
        for run in p.runs:
            if is_header:
                run.font.name = style.typography.body.name
                run.font.size = Pt(11)
                run.font.bold = True
                run.font.color.rgb = style.text_on_dark
            else:
                run.font.name = style.typography.body.name
                run.font.size = Pt(10)
                if highlight_color:
                    run.font.color.rgb = highlight_color
                    run.font.bold = True
                else:
                    run.font.color.rgb = style.text_primary

    cell.fill.solid()
    if is_header:
        cell.fill.fore_color.rgb = style.primary
    elif is_alt:
        cell.fill.fore_color.rgb = style.background_alt
    else:
        cell.fill.fore_color.rgb = style.background

    cell.margin_left = Inches(0.08)
    cell.margin_right = Inches(0.08)
    cell.margin_top = Inches(0.04)
    cell.margin_bottom = Inches(0.04)


def _get_highlight(val, col_idx, rules, style):
    for rule in rules:
        if rule.get("column") != col_idx:
            continue
        v = str(val).strip()
        if rule["condition"] == "negative" and (v.startswith("-") or v.startswith("(")):
            return style.negative
        if rule["condition"] == "positive" and (v.startswith("+") or (v and v[0].isdigit())):
            return style.positive
    return None


# ---------------------------------------------------------------------------
# Timeline
# ---------------------------------------------------------------------------

def render_timeline(slide, steps, left, top, width, height, style):
    """Render a horizontal timeline with milestone markers.

    Adapts label sizes and spacing for different milestone counts:
    - 1-6: normal layout
    - 7-10: compressed labels, smaller font
    - 11+: should have been split by overflow module
    """
    n = len(steps)
    if n == 0:
        return

    from style_resolver import FontRole

    # Adaptive sizing based on milestone count
    if n <= 6:
        circle_size = 0.30
        label_width = 1.8
        date_font_pt = 12
        label_font_pt = 12
        detail_font_pt = 10
        show_detail = True
    elif n <= 10:
        circle_size = 0.22
        label_width = min(1.4, (width - 1.0) / n * 0.9)
        date_font_pt = 9
        label_font_pt = 9
        detail_font_pt = 8
        show_detail = n <= 8
    else:
        circle_size = 0.18
        label_width = min(1.0, (width - 1.0) / n * 0.85)
        date_font_pt = 8
        label_font_pt = 8
        detail_font_pt = 7
        show_detail = False

    line_top = top + height * 0.40
    line_left = left + 0.50
    line_width = width - 1.00

    # Connector line
    render_panel(slide, line_left, line_top + circle_size / 2 - 0.02, line_width, 0.04, style.text_secondary)

    # Milestones
    spacing = line_width / max(n - 1, 1) if n > 1 else 0
    half_label = label_width / 2

    for i, step in enumerate(steps):
        cx = line_left + i * spacing if n > 1 else line_left + line_width / 2

        # Circle marker
        circle = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(cx - circle_size / 2), Inches(line_top),
            Inches(circle_size), Inches(circle_size),
        )
        circle.fill.solid()
        circle.fill.fore_color.rgb = style.accent
        circle.line.fill.background()

        # Date (above)
        render_text(
            slide, step.get("when", ""),
            cx - half_label, line_top - 0.45, label_width, 0.35,
            FontRole(style.typography.body.name, date_font_pt), style.accent,
            alignment=PP_ALIGN.CENTER,
        )

        # Label (below)
        render_text(
            slide, step.get("what", ""),
            cx - half_label, line_top + circle_size + 0.08, label_width, 0.32,
            FontRole(style.typography.body.name, label_font_pt, bold=True), style.text_primary,
            alignment=PP_ALIGN.CENTER,
        )

        # Detail (further below, only if space allows)
        detail = step.get("detail", "")
        if detail and show_detail:
            render_text(
                slide, detail,
                cx - half_label, line_top + circle_size + 0.40, label_width, 0.45,
                FontRole(style.typography.body.name, detail_font_pt), style.text_secondary,
                alignment=PP_ALIGN.CENTER,
            )


# ---------------------------------------------------------------------------
# Emphasis (big number / quote)
# ---------------------------------------------------------------------------

def render_emphasis(slide, emphasis_data, left, top, width, height, style):
    """Render an emphasis element (big number or quote)."""
    emp_type = emphasis_data.get("type", "number")

    if emp_type == "number":
        from style_resolver import FontRole
        value = emphasis_data.get("value", "")
        label = emphasis_data.get("label", "")
        context = emphasis_data.get("context", "")

        # Big number
        render_text(
            slide, value, left, top + 0.30, width, 1.5,
            FontRole(style.typography.display.name, 60, bold=True), style.accent,
            alignment=PP_ALIGN.CENTER,
        )

        if label:
            render_text(
                slide, label, left, top + 1.90, width, 0.60,
                style.typography.heading, style.text_primary,
                alignment=PP_ALIGN.CENTER,
            )

        if context:
            # Use proportional insets that work for any zone width
            inset = min(1.5, width * 0.12)
            render_text(
                slide, context, left + inset, top + 2.70, max(2.0, width - 2 * inset), 0.80,
                style.typography.body, style.text_secondary,
                alignment=PP_ALIGN.CENTER,
            )

    elif emp_type == "quote":
        from style_resolver import FontRole
        quote_text = emphasis_data.get("text", "")
        attribution = emphasis_data.get("attribution", "")

        # Opening quote mark
        render_text(
            slide, "\u201C", left + 0.30, top, 1.0, 1.0,
            FontRole(style.typography.display.name, 60, bold=True), style.accent,
        )

        # Quote text
        render_text(
            slide, quote_text, left + 1.2, top + 0.40, width - 2.0, height - 1.50,
            FontRole(style.typography.body.name, 20, italic=True), style.text_primary,
        )

        if attribution:
            render_text(
                slide, f"\u2014 {attribution}",
                left + 1.2, top + height - 1.0, width - 2.0, 0.40,
                style.typography.body, style.text_secondary,
                alignment=PP_ALIGN.RIGHT,
            )


# ---------------------------------------------------------------------------
# Evaluation (pros/cons)
# ---------------------------------------------------------------------------

def render_evaluation(slide, options, left, top, width, height, style):
    """Render pros/cons comparison cards."""
    n = len(options)
    if n == 0:
        return

    CARD_GUTTER = 0.40
    card_width = (width - (n - 1) * CARD_GUTTER) / n
    card_height = height - 0.20
    card_padding = 0.25

    for i, option in enumerate(options[:min(n, 4)]):
        col_left = left + i * (card_width + CARD_GUTTER)
        heading = option.get("heading", "")
        inner_left = col_left + card_padding
        inner_width = card_width - 2 * card_padding

        # Card background
        render_rounded_card(slide, col_left, top, card_width, card_height, style.background_alt)

        y = top + 0.15
        if heading:
            # Use body font for heading if card is narrow
            heading_font = style.typography.heading if card_width > 4.0 else style.typography.body
            render_text(
                slide, heading, inner_left, y, inner_width, 0.40,
                heading_font, style.primary,
            )
            render_accent_bar(slide, inner_left, y + 0.43, inner_width, style)
            y += 0.60

        # Advantages
        pros = option.get("advantages", option.get("pros", []))
        if pros:
            from style_resolver import FontRole
            label_font = FontRole(style.typography.body.name, 10, bold=True)
            item_font = style.typography.body_small
            render_text(
                slide, "Advantages", inner_left, y, inner_width, 0.22,
                label_font, style.positive,
            )
            y += 0.24
            for item in pros:
                render_text(
                    slide, f"\u2713  {item}", inner_left, y, inner_width, 0.24,
                    item_font, style.text_primary,
                )
                y += 0.24

        # Challenges
        cons = option.get("challenges", option.get("cons", []))
        if cons:
            from style_resolver import FontRole
            label_font = FontRole(style.typography.body.name, 10, bold=True)
            item_font = style.typography.body_small
            y += 0.12
            render_text(
                slide, "Challenges", inner_left, y, inner_width, 0.22,
                label_font, style.negative,
            )
            y += 0.24
            for item in cons:
                render_text(
                    slide, f"\u2717  {item}", inner_left, y, inner_width, 0.24,
                    item_font, style.text_primary,
                )
                y += 0.24


# ---------------------------------------------------------------------------
# Numbered items (outline, takeaway)
# ---------------------------------------------------------------------------

def render_numbered_items(slide, items, left, top, width, height, style):
    """Render numbered items with circle markers.

    Calculates spacing based on whether items have detail text.
    """
    has_details = any(
        (isinstance(item, dict) and (item.get("detail") or item.get("subtext")))
        for item in items
    )
    # Items with details need more space
    min_item_height = 0.80 if has_details else 0.50
    item_height = max(min_item_height, height / max(len(items), 1))
    # But don't exceed available space
    if item_height * len(items) > height:
        item_height = height / max(len(items), 1)

    for i, item in enumerate(items):
        if isinstance(item, str):
            text = item
            detail = ""
        else:
            text = item.get("text", item.get("what", ""))
            detail = item.get("detail", item.get("subtext", ""))

        y = top + i * item_height
        circle_size = 0.38

        # Number circle
        circle = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(left), Inches(y + 0.05),
            Inches(circle_size), Inches(circle_size),
        )
        circle.fill.solid()
        circle.fill.fore_color.rgb = style.accent
        circle.line.fill.background()
        tf = circle.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = str(i + 1)
        from style_resolver import FontRole
        _apply_font(run, FontRole(style.typography.body.name, 14, bold=True), style.text_on_dark)

        # Text -- use body font instead of heading for long text
        text_left = left + circle_size + 0.20
        text_width = width - circle_size - 0.20
        text_font = style.typography.body if len(text) > 50 else style.typography.heading
        render_text(
            slide, text, text_left, y + 0.02, text_width, 0.38,
            text_font, style.text_primary,
        )
        if detail:
            render_text(
                slide, detail, text_left, y + 0.40, text_width, 0.30,
                style.typography.body_small, style.text_secondary,
            )


def render_takeaway(slide, data, left, top, width, height, style):
    """Render takeaway points + optional call-to-action."""
    takeaways = data.get("takeaways", [])
    action = data.get("action", "")

    if takeaways:
        list_height = height - (0.70 if action else 0)
        render_numbered_items(slide, takeaways, left, top, width, list_height, style)

    if action:
        cta_top = top + height - 0.60
        render_rounded_card(slide, left, cta_top, width, 0.55, style.accent)
        from style_resolver import FontRole
        render_text(
            slide, action, left + 0.20, cta_top + 0.10, width - 0.40, 0.35,
            FontRole(style.typography.body.name, 14, bold=True), style.text_on_dark,
            alignment=PP_ALIGN.CENTER,
        )


# ---------------------------------------------------------------------------
# Footer
# ---------------------------------------------------------------------------

def render_footer(slide, metadata, page_num, total_pages, style):
    """Render footer with separator, company name, and page number.

    Uses fixed absolute positioning to avoid overlap issues from
    variable library margins.
    """
    sw = style.slide_width_inches
    sh = style.slide_height_inches

    # Fixed footer positioning
    FOOTER_LEFT = 0.75
    FOOTER_RIGHT_MARGIN = 0.75
    FOOTER_TOP = sh - 0.55
    FOOTER_WIDTH = sw - FOOTER_LEFT - FOOTER_RIGHT_MARGIN

    render_separator_line(slide, FOOTER_LEFT, FOOTER_TOP - 0.05, FOOTER_WIDTH, style)

    company = metadata.get("company", "")
    if company:
        render_text(
            slide, company, FOOTER_LEFT, FOOTER_TOP, 3.0, 0.30,
            style.typography.footer, style.text_secondary,
        )

    confidentiality = metadata.get("confidentiality", "")
    if confidentiality:
        render_text(
            slide, confidentiality,
            FOOTER_LEFT + FOOTER_WIDTH / 2 - 1.5, FOOTER_TOP, 3.0, 0.30,
            style.typography.footer, style.text_secondary,
            alignment=PP_ALIGN.CENTER,
        )

    render_text(
        slide, f"{page_num} / {total_pages}",
        FOOTER_LEFT + FOOTER_WIDTH - 2.0, FOOTER_TOP, 2.0, 0.30,
        style.typography.footer, style.text_secondary,
        alignment=PP_ALIGN.RIGHT,
    )
