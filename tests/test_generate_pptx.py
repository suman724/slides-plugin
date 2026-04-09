"""End-to-end tests for PPTX generation."""

import sys
import os
import json
import tempfile
import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'skills', 'slides', 'scripts'))

from pptx import Presentation
from generate_pptx import generate, validate_spec
from design_system import load_theme, Grid, THEMES
from layouts import LAYOUT_REGISTRY


class TestValidation:
    def test_valid_spec(self):
        spec = {
            "metadata": {"title": "Test"},
            "slides": [{"type": "content", "title": "Slide", "body": []}],
        }
        assert validate_spec(spec) == []

    def test_missing_metadata(self):
        spec = {"slides": [{"type": "content", "title": "X", "body": []}]}
        errors = validate_spec(spec)
        assert any("metadata" in e for e in errors)

    def test_missing_slides(self):
        spec = {"metadata": {"title": "Test"}}
        errors = validate_spec(spec)
        assert any("slides" in e for e in errors)

    def test_empty_slides(self):
        spec = {"metadata": {"title": "Test"}, "slides": []}
        errors = validate_spec(spec)
        assert any("empty" in e for e in errors)

    def test_unknown_slide_type(self):
        spec = {
            "metadata": {"title": "Test"},
            "slides": [{"type": "nonexistent"}],
        }
        errors = validate_spec(spec)
        assert any("nonexistent" in e for e in errors)

    def test_missing_type_field(self):
        spec = {
            "metadata": {"title": "Test"},
            "slides": [{"title": "No type"}],
        }
        errors = validate_spec(spec)
        assert any("type" in e for e in errors)


class TestDesignSystem:
    def test_all_themes_load(self):
        for name in ["corporate_blue", "modern_dark", "minimal"]:
            theme = load_theme(name)
            assert theme.name == name

    def test_unknown_theme_defaults(self):
        theme = load_theme("nonexistent")
        assert theme.name == "corporate_blue"

    def test_grid_content_width(self):
        grid = Grid()
        expected = grid.slide_width - grid.margin_left - grid.margin_right
        assert abs(grid.content_width - expected) < 0.001

    def test_grid_column_widths(self):
        grid = Grid()
        for n in [1, 2, 3, 4]:
            col_w = grid.col_width(n)
            total = col_w * n + grid.gutter * (n - 1)
            assert abs(total - grid.content_width) < 0.001

    def test_grid_col_positions(self):
        grid = Grid()
        assert grid.col_left(0, 2) == grid.margin_left
        col1_left = grid.col_left(1, 2)
        assert col1_left > grid.margin_left


class TestLayoutRegistry:
    def test_all_layouts_registered(self):
        expected = [
            "title", "closing", "agenda", "section_divider",
            "content", "two_column", "three_column",
            "chart", "table", "key_metrics",
            "comparison", "timeline", "takeaway",
            "quote", "image_with_text", "big_number",
        ]
        for layout in expected:
            assert layout in LAYOUT_REGISTRY, f"Missing layout: {layout}"

    def test_all_layouts_callable(self):
        for name, fn in LAYOUT_REGISTRY.items():
            assert callable(fn), f"Layout {name} is not callable"


class TestGeneration:
    def test_minimal_presentation(self):
        spec = {
            "metadata": {"title": "Test", "theme": "corporate_blue"},
            "slides": [
                {"type": "title", "title": "Test Deck", "subtitle": "Testing"},
                {"type": "content", "title": "Content", "body": [
                    {"type": "bullet", "text": "Point 1", "level": 0},
                ]},
                {"type": "closing", "title": "Thanks"},
            ],
        }
        prs = generate(spec)
        assert len(prs.slides) == 3

    def test_all_slide_types(self):
        """Generate a presentation with every slide type."""
        spec_path = os.path.join(os.path.dirname(__file__), "sample_spec.json")
        with open(spec_path) as f:
            spec = json.load(f)

        prs = generate(spec)
        assert len(prs.slides) == len(spec["slides"])
        # Every slide should have at least 1 shape
        for i, slide in enumerate(prs.slides):
            assert len(slide.shapes) > 0, f"Slide {i+1} has no shapes"

    def test_each_theme(self):
        """Each theme should produce valid output."""
        for theme_name in THEMES:
            spec = {
                "metadata": {"title": "Theme Test", "theme": theme_name},
                "slides": [
                    {"type": "title", "title": f"Theme: {theme_name}"},
                    {"type": "content", "title": "Content", "body": [
                        {"type": "bullet", "text": "Test", "level": 0},
                    ]},
                ],
            }
            prs = generate(spec)
            assert len(prs.slides) == 2

    def test_save_to_file(self):
        spec = {
            "metadata": {"title": "Save Test"},
            "slides": [{"type": "title", "title": "Test"}],
        }
        prs = generate(spec)
        with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as f:
            prs.save(f.name)
            assert os.path.getsize(f.name) > 0
            os.unlink(f.name)

    def test_chart_slide(self):
        spec = {
            "metadata": {"title": "Chart Test"},
            "slides": [{
                "type": "chart",
                "title": "Revenue",
                "chart_type": "column_clustered",
                "categories": ["Q1", "Q2", "Q3"],
                "series": [{"name": "Rev", "values": [10, 20, 30]}],
            }],
        }
        prs = generate(spec)
        assert len(prs.slides) == 1

    def test_table_slide(self):
        spec = {
            "metadata": {"title": "Table Test"},
            "slides": [{
                "type": "table",
                "title": "Data",
                "columns": ["A", "B", "C"],
                "rows": [["1", "2", "3"], ["4", "5", "6"]],
            }],
        }
        prs = generate(spec)
        assert len(prs.slides) == 1

    def test_key_metrics_slide(self):
        spec = {
            "metadata": {"title": "KPI Test"},
            "slides": [{
                "type": "key_metrics",
                "title": "KPIs",
                "metrics": [
                    {"label": "ARR", "value": "$100M", "delta": "+10%", "direction": "up"},
                    {"label": "NRR", "value": "110%", "delta": "+5pp", "direction": "up"},
                    {"label": "Churn", "value": "5%", "delta": "+1pp", "direction": "down"},
                ],
            }],
        }
        prs = generate(spec)
        assert len(prs.slides) == 1

    def test_slide_dimensions(self):
        spec = {
            "metadata": {"title": "Dims"},
            "slides": [{"type": "title", "title": "Test"}],
        }
        prs = generate(spec)
        grid = Grid()
        # Check 16:9 dimensions
        width_in = prs.slide_width / 914400
        height_in = prs.slide_height / 914400
        assert abs(width_in - grid.slide_width) < 0.01
        assert abs(height_in - grid.slide_height) < 0.01
